"""Step 14: which modifiers win among ALL co-op games on Steam.

  python 14_coop_market.py collect   # app ids via store search, details via IStoreBrowseService
  python 14_coop_market.py analyze   # tag and direction uplift -> data/coop_tags.csv, data/coop_directions.csv
  python 14_coop_market.py near      # same directions inside co-op games about animals / nature / hunting

Base: co-op games (tags Co-op / Online Co-Op / Local Co-Op), released 2022-01-01 .. 12 months
before the data date (so every game had a year to collect reviews), paid, not AAA.
For each tag: share of games reaching 100+ (and 1000+) reviews WITH the tag vs WITHOUT it.
A tag counts only if it is among the game's top 10 tags, i.e. the game is really about it.
"""
import argparse
import json
import re
import time

import pandas as pd
import requests
from bs4 import BeautifulSoup

from config import AAA_PUBLISHERS, DATA, RAW

SEARCH_TAGS = {"Co-op": 1685, "Online Co-Op": 3843, "Local Co-Op": 3841}
IDS = RAW / "coop_appids.csv"
ITEMS = RAW / "coop_items.jsonl"
DATA_DATE = pd.Timestamp("2026-10-04")
START = pd.Timestamp("2022-01-01")
MATURE_END = DATA_DATE - pd.DateOffset(months=12)
TOP_TAGS = 10
MIN_GAMES = 40

# directions = groups of tags, named the way a team would talk about them
DIRECTIONS = {
    "Хоррор": ["Horror", "Psychological Horror", "Survival Horror"],
    "Шутер": ["Shooter", "FPS", "Third-Person Shooter", "Looter Shooter", "Arena Shooter"],
    "Экшен / взрывы": ["Action", "Explosions", "Hack and Slash", "Beat 'em up", "Bullet Hell", "Twin Stick Shooter"],
    "Смешное / физика": ["Funny", "Comedy", "Physics", "Memes", "Dark Comedy"],
    "Пати-игра": ["Party Game", "Party"],
    "Выживание / крафт": ["Survival", "Crafting", "Open World Survival Craft", "Base Building"],
    "Рогалик": ["Roguelite", "Roguelike", "Action Roguelike", "Roguelike Deckbuilder"],
    "Экстракшн": ["Extraction Shooter"],
    "Симулятор": ["Simulation", "Life Sim", "Job Simulator", "Immersive Sim"],
    "Уютное / милое": ["Cozy", "Relaxing", "Cute", "Wholesome", "Family Friendly"],
    "Песочница / строительство": ["Sandbox", "Building", "Automation", "City Builder"],
    "Открытый мир / исследование": ["Open World", "Exploration"],
    "Головоломки": ["Puzzle", "Puzzle Platformer", "Logic"],
    "Платформер": ["Platformer", "2D Platformer", "3D Platformer", "Precision Platformer"],
    "PvP / соревновательное": ["PvP", "Competitive", "Team-Based", "Battle Royale", "eSports"],
    "Стратегия / тактика": ["Strategy", "Tactical", "Tower Defense", "RTS", "Turn-Based Strategy"],
    "RPG": ["RPG", "Action RPG", "Loot", "Character Customization"],
    "Стелс": ["Stealth"],
    "Гонки / транспорт": ["Racing", "Driving", "Vehicular Combat", "Automobile Sim"],
    "Спорт": ["Sports", "Football (Soccer)", "Basketball", "Golf"],
    "Животные / природа": ["Animals", "Nature", "Dogs", "Cats", "Dinosaurs", "Hunting", "Fishing"],
    "Мультяшное / стилизация": ["Cartoony", "Stylized", "Colorful", "Hand-drawn", "Cartoon"],
    "Реалистичное": ["Realistic"],
    "Процедурная генерация": ["Procedural Generation"],
    "Волны / орды": ["Wave Shooter", "Zombies", "Horde"],
}

_s = requests.Session()
_s.headers["User-Agent"] = "Mozilla/5.0 (market research script)"


def get(url, **params):
    for attempt in range(8):
        try:
            r = _s.get(url, params=params, timeout=60)
            if r.status_code == 200:
                return r
        except requests.RequestException:
            pass
        time.sleep(min(60, 2 ** attempt))
    return None


def search_ids(tagid):
    ids, start, total = set(), 0, None
    while total is None or start < total:
        r = get("https://store.steampowered.com/search/results/", tags=tagid, category1=998,
                infinite=1, start=start, count=100, cc="us", l="english")
        if r is None:
            break
        d = r.json()
        total = d["total_count"]
        rows = BeautifulSoup(d["results_html"], "html.parser").select("a.search_result_row")
        for a in rows:
            v = a.get("data-ds-appid", "")
            if v and "," not in v:
                ids.add(int(v))
        if not rows:
            break
        start += 100
        time.sleep(0.7)
    return ids, total


def cmd_collect(args):
    if not IDS.exists():
        allids = set()
        for name, tid in SEARCH_TAGS.items():
            ids, total = search_ids(tid)
            print(f"{name}: total_count={total} collected={len(ids)}", flush=True)
            allids |= ids
        pd.DataFrame({"appid": sorted(allids)}).to_csv(IDS, index=False)
    ids = pd.read_csv(IDS).appid.tolist()
    done = set()
    if ITEMS.exists():
        done = {json.loads(l)["appid"] for l in ITEMS.open()}
    todo = [a for a in ids if a not in done]
    print(f"apps: {len(ids)}  to fetch: {len(todo)}", flush=True)
    names = {t["tagid"]: t["name"] for t in
             get("https://store.steampowered.com/tagdata/populartags/english").json()}
    with ITEMS.open("a") as f:
        for i in range(0, len(todo), 100):
            chunk = todo[i:i + 100]
            req = {"ids": [{"appid": a} for a in chunk],
                   "context": {"language": "english", "country_code": "US"},
                   "data_request": {"include_tag_count": 20, "include_basic_info": True,
                                    "include_release": True, "include_reviews": True}}
            r = get("https://api.steampowered.com/IStoreBrowseService/GetItems/v1/",
                    input_json=json.dumps(req))
            items = r.json().get("response", {}).get("store_items", []) if r is not None else []
            for it in items:
                tags = [names.get(t["tagid"], str(t["tagid"]))
                        for t in sorted(it.get("tags", []), key=lambda t: -t.get("weight", 0))]
                rel = it.get("release", {})
                rv = it.get("reviews", {}).get("summary_filtered", {})
                bi = it.get("basic_info", {})
                f.write(json.dumps({
                    "appid": it.get("appid", it.get("id")), "name": it.get("name", ""),
                    "type": it.get("type"), "visible": it.get("visible"),
                    "is_free": it.get("is_free", False),
                    "early_access": it.get("is_early_access", False),
                    "coming_soon": it.get("is_coming_soon", False),
                    "release_ts": rel.get("steam_release_date") or rel.get("original_release_date"),
                    "reviews": rv.get("review_count", 0), "positive_pct": rv.get("percent_positive"),
                    "publishers": [p["name"] for p in bi.get("publishers", [])],
                    "developers": [p["name"] for p in bi.get("developers", [])],
                    "tags": tags,
                }, ensure_ascii=False) + "\n")
            if (i // 100) % 20 == 0:
                print(f"{i + len(chunk)}/{len(todo)}", flush=True)
            time.sleep(0.5)
    print("finished", flush=True)


def load_base():
    rows = [json.loads(l) for l in ITEMS.open()]
    df = pd.DataFrame(rows).drop_duplicates("appid")
    df = df[df.type == 0]  # games only
    df["date"] = pd.to_datetime(df.release_ts, unit="s", errors="coerce")
    pubs = df.publishers.apply(lambda p: " | ".join(p).lower()) + " | " + \
        df.developers.apply(lambda p: " | ".join(p).lower())
    aaa = r"(?<![\w])(?:" + "|".join(re.escape(p) for p in AAA_PUBLISHERS) + r")(?![\w])"
    df["is_aaa"] = pubs.str.contains(aaa, regex=True)
    df["top_tags"] = df.tags.apply(lambda t: set(t[:TOP_TAGS]))
    df["all_tags"] = df.tags.apply(set)
    base = df[~df.coming_soon & ~df.is_free & ~df.is_aaa & df.date.notna()
              & (df.date >= START) & (df.date <= MATURE_END)].copy()
    base["period"] = (base.date.dt.year <= 2023).map({True: "2022–23", False: "2024–25"})
    return df, base


def rate(g, col="reviews", thr=100):
    return round(100 * (g[col] >= thr).mean(), 1) if len(g) else None


def uplift_rows(base, label, has):
    w, wo = base[has], base[~has]
    early, late = w[w.period == "2022–23"], w[w.period == "2024–25"]
    return {
        "name": label, "games": len(w),
        "pct_100plus": rate(w), "pct_100plus_without": rate(wo),
        "uplift_100": round(rate(w) / rate(wo), 2) if len(w) and rate(wo) else None,
        "pct_1000plus": rate(w, thr=1000), "pct_1000plus_without": rate(wo, thr=1000),
        "uplift_1000": round(rate(w, thr=1000) / rate(wo, thr=1000), 2) if len(w) and rate(wo, thr=1000) else None,
        "median_reviews": float(w.reviews.median()) if len(w) else None,
        "pct_100plus_2022_23": rate(early), "games_2022_23": len(early),
        "pct_100plus_2024_25": rate(late), "games_2024_25": len(late),
        "share_of_coop_games_2024_25": round(100 * len(late) / max(1, (base.period == "2024–25").sum()), 1),
        "examples_top": "; ".join(f"{n} ({r})" for n, r in
                                  w.sort_values("reviews", ascending=False).head(3)[["name", "reviews"]].values),
    }


def cmd_analyze(args):
    df, base = load_base()
    print(f"co-op apps: {len(df)}  base (released {START.date()}..{MATURE_END.date()}, paid, non-AAA): {len(base)}")
    print(f"base: 100+ = {rate(base)}%, 1000+ = {rate(base, thr=1000)}%")
    tag_counts = pd.Series([t for s in base.top_tags for t in s]).value_counts()
    rows = [uplift_rows(base, t, base.top_tags.apply(lambda s, t=t: t in s))
            for t, n in tag_counts.items() if n >= MIN_GAMES]
    tags = pd.DataFrame(rows).sort_values("pct_100plus", ascending=False)
    tags.to_csv(DATA / "coop_tags.csv", index=False)
    drows = []
    for d, tl in DIRECTIONS.items():
        has = base.top_tags.apply(lambda s, tl=set(tl): bool(s & tl))
        r = uplift_rows(base, d, has)
        r["tags"] = ", ".join(tl)
        drows.append(r)
    dirs = pd.DataFrame(drows).sort_values("pct_100plus", ascending=False)
    dirs.to_csv(DATA / "coop_directions.csv", index=False)
    export(base, dirs, tags)
    with pd.option_context("display.width", 250, "display.max_columns", 20, "display.max_colwidth", 60):
        print(dirs.drop(columns=["tags", "examples_top"]).to_string(index=False))


def export(base, dirs, tags):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    cols = {"name": "Направление", "games": "Игр", "pct_100plus": "Набрали 100+, %",
            "pct_100plus_without": "Без него, %", "uplift_100": "Во сколько раз чаще",
            "pct_1000plus": "Набрали 1000+, %", "pct_1000plus_without": "1000+ без него, %",
            "uplift_1000": "1000+: во сколько раз", "median_reviews": "Медиана отзывов",
            "pct_100plus_2022_23": "100+ в 2022–23, %", "pct_100plus_2024_25": "100+ в 2024–25, %",
            "share_of_coop_games_2024_25": "Доля кооп-игр 2024–25, %",
            "examples_top": "Крупнейшие игры", "tags": "Теги"}

    def prep(df, first):
        out = df[[c for c in cols if c in df.columns]].rename(columns={**cols, "name": first})
        out.insert(2, "Надёжность", ["мало игр" if n < MIN_GAMES else "" for n in df.games])
        return out

    legend = pd.DataFrame([
        ("База", f"Все кооп-игры Steam (теги Co-op / Online Co-Op / Local Co-Op), вышедшие "
                 f"{START.date()} – {MATURE_END.date()} (у каждой был год на отзывы), платные, без AAA: "
                 f"{len(base)} игр. В среднем 100+ отзывов набрали {rate(base)}%, 1000+ — {rate(base, thr=1000)}%."),
        ("Как считается", "Игра относится к направлению, если хотя бы один его тег — среди её 10 главных тегов. "
                          "«Во сколько раз чаще» = доля 100+ у игр направления ÷ доля 100+ у остальных кооп-игр."),
        ("Надёжность", f"«мало игр» — меньше {MIN_GAMES} игр, цифра неустойчива."),
        ("Осторожно: масштаб", "Направления вроде симулятора, выживания и открытого мира обычно требуют "
                               "большего бюджета, а платформеры и аркады — меньшего. Часть разницы — это разница "
                               "в масштабе игр, а не только в спросе на направление."),
        ("Полнота", "Поиск Steam отдал ~93% кооп-игр; для долей это не критично."),
    ], columns=["Что", "Пояснение"])
    with pd.ExcelWriter(DATA / "coop_directions.xlsx", engine="openpyxl") as w:
        for name, df, widths in (("Пояснения", legend, {"Что": 20, "Пояснение": 140}),
                                 ("Направления", prep(dirs, "Направление"), {"Направление": 28, "Крупнейшие игры": 70, "Теги": 60}),
                                 ("Все теги", prep(tags, "Тег"), {"Тег": 26, "Крупнейшие игры": 70})):
            df.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for i, c in enumerate(df.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 13)
                ws.cell(1, i).font = Font(bold=True)


THEME = {"Animals", "Nature", "Hunting", "Dogs", "Cats", "Fishing", "Dinosaurs", "Horses",
         "Birds", "Wolves", "Sharks", "Zoo"}
DOG_NAME = re.compile(r"\b(dogs?|pupp(y|ies)|hounds?|doggo|corgis?|retrievers?)\b", re.I)


def wilson(k, n, z=1.96):
    """95% interval for a share, in %. Honest for small n, unlike ± from the normal approximation."""
    if not n:
        return None, None
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * ((p * (1 - p) / n + z * z / (4 * n * n)) ** 0.5) / d
    return round(100 * (c - h), 1), round(100 * (c + h), 1)


def cmd_near(args):
    """Check the co-op direction signal on co-op games about animals, nature and hunting."""
    df, base = load_base()
    themed = base.all_tags.apply(lambda s: bool(s & THEME)) | base.name.str.contains(DOG_NAME)
    near = base[themed]
    full = pd.read_csv(DATA / "coop_directions.csv").set_index("name")
    print(f"near-base: {len(near)} co-op games about animals/nature/hunting (of {len(base)}); "
          f"100+ = {rate(near)}%, 1000+ = {rate(near, thr=1000)}%")
    rows = []
    for d, tl in DIRECTIONS.items():
        if d == "Животные / природа":
            continue  # that is the theme itself
        has = near.top_tags.apply(lambda s, tl=set(tl): bool(s & tl))
        w, wo = near[has], near[~has]
        if len(w) < args.min_games:
            continue
        k = int((w.reviews >= 100).sum())
        lo, hi = wilson(k, len(w))
        up = round(rate(w) / rate(wo), 2) if rate(wo) else None
        fu = full.loc[d, "uplift_100"] if d in full.index else None
        rows.append({
            "direction": d, "games": len(w), "pct_100plus": rate(w), "ci_low": lo, "ci_high": hi,
            "pct_100plus_without": rate(wo), "uplift_100": up,
            "uplift_100_all_coop": fu,
            "agrees_with_all_coop": None if up is None or fu is None else
            ("да" if (up - 1) * (fu - 1) > 0 or abs(up - 1) < 0.1 and abs(fu - 1) < 0.1 else "нет"),
            "pct_1000plus": rate(w, thr=1000),
            "examples_top": "; ".join(f"{n} ({r})" for n, r in
                                      w.sort_values("reviews", ascending=False).head(4)[["name", "reviews"]].values),
        })
    out = pd.DataFrame(rows).sort_values("pct_100plus", ascending=False)
    out.to_csv(DATA / "coop_near_directions.csv", index=False)
    near.sort_values("reviews", ascending=False)[["appid", "name", "reviews", "positive_pct", "date"]].assign(
        tags=near.tags.apply(lambda t: ", ".join(t[:10]))).to_csv(DATA / "coop_near_games.csv", index=False)
    export_near(near, base, out)
    with pd.option_context("display.width", 250, "display.max_columns", 20, "display.max_colwidth", 80):
        print(out.to_string(index=False))


def export_near(near, base, out):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    legend = pd.DataFrame([
        ("Выборка", f"Кооп-игры из общей базы ({len(base)} игр, 2022 – {MATURE_END.date()}, платные, без AAA), "
                    f"у которых среди тегов есть животные, природа или охота ({', '.join(sorted(THEME))}), "
                    f"или собака в названии: {len(near)} игр. 100+ набрали {rate(near)}%, 1000+ — {rate(near, thr=1000)}%."),
        ("Интервал", "95% интервал для доли 100+ (метод Уилсона). Широкий интервал = мало игр, цифра неустойчива."),
        ("Совпадает с коопом", "«да», если направление и здесь, и в общем коопе действует в одну сторону "
                               "(выше или ниже остальных игр)."),
        ("Спорт", "В «спорт» здесь попали симуляторы охоты и рыбалки (Way of the Hunter, The Angler) — это не спорт в обычном смысле."),
    ], columns=["Что", "Пояснение"])
    cols = {"direction": "Направление", "games": "Игр", "pct_100plus": "Набрали 100+, %",
            "ci_low": "Интервал от, %", "ci_high": "Интервал до, %", "pct_100plus_without": "Без него, %",
            "uplift_100": "Во сколько раз чаще", "uplift_100_all_coop": "То же во всём коопе",
            "agrees_with_all_coop": "Совпадает с коопом", "pct_1000plus": "Набрали 1000+, %",
            "examples_top": "Крупнейшие игры"}
    games = near.sort_values("reviews", ascending=False).assign(
        year=near.date.dt.year, tags10=near.tags.apply(lambda t: ", ".join(t[:10])),
        url="https://store.steampowered.com/app/" + near.appid.astype(str) + "/")[
        ["name", "year", "reviews", "positive_pct", "tags10", "url"]].rename(columns={
            "name": "Игра", "year": "Год", "reviews": "Отзывы", "positive_pct": "% положит.",
            "tags10": "Главные теги", "url": "Ссылка"})
    with pd.ExcelWriter(DATA / "coop_near.xlsx", engine="openpyxl") as w:
        for name, df, widths in (("Пояснения", legend, {"Что": 20, "Пояснение": 140}),
                                 ("Направления", out.rename(columns=cols), {"Направление": 28, "Крупнейшие игры": 80}),
                                 ("Игры", games, {"Игра": 40, "Главные теги": 90, "Ссылка": 45})):
            df.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for i, c in enumerate(df.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 13)
                ws.cell(1, i).font = Font(bold=True)


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("collect")
    sub.add_parser("analyze")
    n = sub.add_parser("near")
    n.add_argument("--min-games", type=int, default=10)
    args = p.parse_args()
    {"collect": cmd_collect, "analyze": cmd_analyze, "near": cmd_near}[args.cmd](args)


if __name__ == "__main__":
    main()
