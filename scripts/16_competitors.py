"""Step 16: announced (not yet released) competitors.

  python 16_competitors.py

Tiers:
  A — direct: dog + hunting, dog + co-op, hunting + co-op (from the dog/hunting dataset);
  B — theme + format: online co-op about animals / nature / hunting with stealth, survival or horror;
  C — format only: online co-op with stealth among the top tags.
Release dates and descriptions are refreshed through IStoreBrowseService.
Output: data/competitors.csv, data/competitors.xlsx
"""
import json
import re
import time

import pandas as pd
import requests
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import AAA_PUBLISHERS, DATA, RAW

THEME = {"Animals", "Nature", "Hunting", "Dogs", "Wolves", "Fishing"}
TENSION = {"Stealth", "Survival", "Horror", "Psychological Horror", "Survival Horror"}
COOP = {"Co-op", "Online Co-Op", "Local Co-Op", "Multiplayer"}
DATA_DATE = pd.Timestamp("2026-10-04")
# picked by hand from tiers A–C after reading the descriptions: closest to "dog + hunting + co-op"
CLOSEST = {
    "PAWTENTIAL DISASTER": "хаотичный физический кооп за стаю собак",
    "Salty Dogs": "кооп-отряд вооружённых собак, 2D-экшен",
    "Quack Hunters": "физический кооп-симулятор охоты на 1–4 игроков с юмором и хоррором",
    "How to Hunt": "физический кооп-симулятор охоты на 1–4 игроков",
    "Hunting Simulator 3": "реалистичная охота с собакой, открытый мир",
    "Man and Dog": "симулятор охоты с собакой, онлайн-кооп",
    "Ultimate Hunting®": "реалистичная охота, есть собака",
    "Blue Ridge Hunting": "кооп-хоррор: охота на криптидов",
    "Meat Safari": "кооп: охота на ингредиенты + готовка, пати",
    "Fearless Crew": "кооп-охота на монстров с выносом добычи на корабль",
    "Barely Alive": "кооп-выживание, пати: «охоться на друзей»",
    "Flow's Island": "кооп-выживание за бродячих зверей",
    "Dirty Dogs": "социальная дедукция за собак (пати)",
    "Hide or Seek: Container Island": "прятки: кошки против собак, кооп/PvP",
    "Fading Day": "кооп-хоррор выживание, есть собака",
}


def get_items(ids):
    names = {t["tagid"]: t["name"] for t in
             requests.get("https://store.steampowered.com/tagdata/populartags/english", timeout=60).json()}
    out = {}
    for i in range(0, len(ids), 100):
        req = {"ids": [{"appid": int(a)} for a in ids[i:i + 100]],
               "context": {"language": "english", "country_code": "US"},
               "data_request": {"include_basic_info": True, "include_release": True, "include_tag_count": 20}}
        for attempt in range(6):
            r = requests.get("https://api.steampowered.com/IStoreBrowseService/GetItems/v1/",
                             params={"input_json": json.dumps(req)}, timeout=60)
            if r.status_code == 200:
                break
            time.sleep(2 ** attempt)
        for it in r.json().get("response", {}).get("store_items", []):
            rel = it.get("release", {})
            ts = rel.get("steam_release_date")
            out[it.get("appid", it.get("id"))] = {
                "coming_soon": it.get("is_coming_soon", False),
                "release": (pd.to_datetime(ts, unit="s").strftime("%Y-%m-%d") if ts
                            else rel.get("custom_release_date_message") or rel.get("coming_soon_display", "")),
                "release_ts": ts,
                "short_description": it.get("basic_info", {}).get("short_description", ""),
                "developers": "; ".join(p["name"] for p in it.get("basic_info", {}).get("developers", [])),
                "publishers": "; ".join(p["name"] for p in it.get("basic_info", {}).get("publishers", [])),
                "tags_api": [names.get(t["tagid"], "") for t in
                             sorted(it.get("tags", []), key=lambda t: -t.get("weight", 0))],
                "name_api": it.get("name", ""),
            }
        time.sleep(0.5)
    return out


def main():
    g = pd.read_csv(DATA / "games.csv")
    for c in [c for c in g.columns if c.startswith("el_")] + ["coming_soon", "is_aaa"]:
        g[c] = g[c].astype(bool)
    up = g[g.coming_soon & ~g.is_aaa].copy()
    up["tag_list"] = up.tags.fillna("").apply(lambda t: [x.strip() for x in t.split(",") if x.strip()])
    coop = up.tag_list.apply(lambda t: bool(set(t) & COOP))
    tier = {}
    for a, d, h, c in zip(up.appid, up.el_dog, up.el_hunting, coop):
        if (d and h) or (d and c) or (h and c):
            tier[a] = "A"
    items = pd.DataFrame([json.loads(l) for l in (RAW / "coop_items.jsonl").open()])
    ci = items[items.coming_soon.astype(bool)]
    for a, t in zip(ci.appid, ci.tags):
        top = set(t[:10])
        if a in tier:
            continue
        if "Online Co-Op" in t and set(t) & THEME and top & TENSION:
            tier[a] = "B"
        elif "Online Co-Op" in top and "Stealth" in top:
            tier[a] = "C"
    info = get_items(list(tier))
    aaa = re.compile(r"(?<![\w])(?:" + "|".join(re.escape(p) for p in AAA_PUBLISHERS) + r")(?![\w])", re.I)
    rows = []
    for a, t in tier.items():
        i = info.get(a)
        if not i or not i["coming_soon"]:
            continue  # released since, or page gone
        if aaa.search(i["developers"] + " | " + i["publishers"]):
            continue
        tags = i["tags_api"]
        dated = i["release_ts"] is not None or bool(re.search(r"20\d\d", str(i["release"])))
        rows.append({"tier": t, "appid": a, "name": i["name_api"], "release": i["release"],
                     "has_date": dated, "developers": i["developers"], "publishers": i["publishers"],
                     "dog": "Dogs" in tags or bool(re.search(r"\bdogs?\b|pupp|hound", i["name_api"], re.I)),
                     "hunting": "Hunting" in tags, "online_coop": "Online Co-Op" in tags,
                     "stealth": "Stealth" in tags, "horror": bool(set(tags) & {"Horror", "Psychological Horror", "Survival Horror"}),
                     "tags": ", ".join(tags[:10]), "short_description": i["short_description"],
                     "url": f"https://store.steampowered.com/app/{a}/"})
    df = pd.DataFrame(rows).sort_values(["tier", "has_date", "release"], ascending=[True, False, True])
    df["why_close"] = df.name.map(CLOSEST)
    df.to_csv(DATA / "competitors.csv", index=False)
    print(df.groupby("tier").agg(games=("appid", "size"), dated=("has_date", "sum")))

    ru = {"tier": "Уровень", "name": "Игра", "release": "Выход", "developers": "Разработчик",
          "publishers": "Издатель", "dog": "Собака", "hunting": "Охота", "online_coop": "Онлайн-кооп",
          "stealth": "Стелс", "horror": "Хоррор", "tags": "Главные теги",
          "short_description": "Описание", "url": "Ссылка", "why_close": "Чем похожа"}
    legend = pd.DataFrame([
        ("A — прямые", "Анонсы с собакой и охотой, или собакой/охотой и коопом."),
        ("B — тема + формат", "Онлайн-кооп про животных, природу или охоту со стелсом, выживанием или хоррором."),
        ("C — формат", "Онлайн-кооп со стелсом среди главных тегов (будущие R.E.P.O.-подобные)."),
        ("Ближайшие", "15 анонсов, отобранных вручную по описаниям как самые близкие к «собака + охота + кооп»."),
        ("Даты", "Анонсы без конкретной даты («To be announced», «Coming soon») часто заброшены — "
                 "они в конце каждого листа."),
        ("Данные", f"Steam на {DATA_DATE.date()}. Без AAA. Вишлисты и подписчики Steam не отдаёт."),
    ], columns=["Что", "Пояснение"])
    with pd.ExcelWriter(DATA / "competitors.xlsx", engine="openpyxl") as w:
        sheets = [("Пояснения", legend)] + [
            (name, df[df.tier == t].drop(columns=["tier", "appid", "has_date"]))
            for t, name in (("A", "A — прямые"), ("B", "B — тема + формат"), ("C", "C — формат"))]
        sheets.insert(1, ("Ближайшие", df[df.why_close.notna()].drop(columns=["appid", "has_date"])))
        for name, d in sheets:
            d = d.rename(columns=ru)
            for c in ("Собака", "Охота", "Онлайн-кооп", "Стелс", "Хоррор"):
                if c in d:
                    d[c] = d[c].map({True: "да", False: ""})
            d.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            widths = {"Игра": 36, "Описание": 90, "Главные теги": 60, "Ссылка": 45, "Пояснение": 120, "Что": 20,
                      "Разработчик": 22, "Издатель": 22, "Выход": 16, "Чем похожа": 50}
            for i, c in enumerate(d.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 10)
                ws.cell(1, i).font = Font(bold=True)


if __name__ == "__main__":
    main()
