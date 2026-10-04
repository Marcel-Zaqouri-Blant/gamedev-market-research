"""Step 9: Excel workbook with audience overlap and creators per intersection.

Output: data/audience_and_creators.xlsx
"""
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import DATA

SEG_RU = {"dog": "собака", "hunting": "охота", "friendslop": "френдслоп",
          "horror": "хоррор", "roguelite": "рогалик", "extraction": "экстракшн"}
SIDE_RU = {"a": "сторона 1", "b": "сторона 2", "hybrid": "гибрид"}


def ru_pair(s):
    return " + ".join(SEG_RU.get(x, x) for x in s.split("+"))


def write(writer, name, df, widths):
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
        ws.cell(1, i).font = Font(bold=True)
    for col_name in ("Канал", "Лучшее видео"):
        if col_name not in df.columns:
            continue
        ci = list(df.columns).index(col_name) + 1
        links = df["_url"] if col_name == "Канал" else df[col_name]
        for r, url in enumerate(links, 2):
            if isinstance(url, str) and url.startswith("http"):
                ws.cell(r, ci).hyperlink, ws.cell(r, ci).style = url, "Hyperlink"
    if "_url" in df.columns:  # helper column: hide it
        ws.column_dimensions[get_column_letter(list(df.columns).index("_url") + 1)].hidden = True


def main():
    pairs = pd.read_csv(DATA / "audience_overlap.csv")
    p = pd.DataFrame({
        "Сегмент 1": pairs.segment_a.map(SEG_RU), "Сегмент 2": pairs.segment_b.map(SEG_RU),
        "Lift": pairs.lift, "Lift от": pairs.lift_low, "Lift до": pairs.lift_high,
        "% аудитории 1, игравшей в 2": pairs.pct_of_a_also_b,
        "% аудитории 2, игравшей в 1": pairs.pct_of_b_also_a,
        "Норма для 2, %": pairs.baseline_pct_b, "Норма для 1, %": pairs.baseline_pct_a,
        "Людей на стыке": pairs.overlap, "Игр 1": pairs.games_a, "Игр 2": pairs.games_b,
        "Надёжно": pairs.reliable.map({True: "да", False: "мало данных"}),
    })
    legend = pd.DataFrame([
        ("Lift", "Во сколько раз чаще аудитория сегмента 1 пишет отзывы на игры сегмента 2 "
                 "(и наоборот, среднее), чем авторы отзывов на остальные игры выборки. "
                 "1 = связи нет, 2 = вдвое чаще, <1 = аудитории расходятся."),
        ("Lift от / до", "Примерный 95% интервал."),
        ("Игры-гибриды", "Игры с обоими элементами исключены из сравнения пары, "
                         "иначе они создают пересечение сами."),
        ("Данные", "Авторы отзывов Steam: до 500 последних на игру, 2 854 игры, ~489 тыс. человек. "
                   "Исключены 280 аккаунтов с отзывами на 20+ игр выборки (кураторы, сборщики ключей)."),
        ("Блогеры", "YouTube: по 50 самых просматриваемых видео (категория «Игры») на каждую игру. "
                    "На стыке — канал, снимавший игру-гибрид или игры обеих сторон. "
                    "Официальные каналы и трейлеры исключены."),
    ], columns=["Что", "Пояснение"])

    cr = pd.read_csv(DATA / "creators.csv")
    targets = pd.read_csv(DATA / "creator_targets.csv")

    with pd.ExcelWriter(DATA / "audience_and_creators.xlsx", engine="openpyxl") as w:
        write(w, "Пояснения", legend, {"Что": 18, "Пояснение": 140})
        write(w, "Пересечения", p, {"Сегмент 1": 13, "Сегмент 2": 13})
        for inter, g in cr.groupby("intersection", sort=False):
            g = g.sort_values("subscribers", ascending=False)
            games = g.games_covered.str.replace(r"\[a\]", "[сторона 1]", regex=True) \
                .str.replace(r"\[b\]", "[сторона 2]", regex=True).str.replace("[hybrid]", "[гибрид]", regex=False)
            out = pd.DataFrame({
                "Канал": g.channel.values, "Подписчики": g.subscribers.values,
                "Страна": g.country.fillna("").values, "Игры (сторона стыка)": games.values,
                "Игр": g.n_games.values, "Просмотры на этих играх": g.views_on_these_games.values,
                "Лучшее видео": g.top_video.values, "_url": g.url.values,
            })
            write(w, ru_pair(inter)[:31], out, {"Канал": 30, "Игры (сторона стыка)": 80,
                                                 "Лучшее видео": 45, "Просмотры на этих играх": 16})
        t = targets.assign(intersection=targets.intersection.map(ru_pair), side=targets.side.map(SIDE_RU))
        t = t.rename(columns={"intersection": "Стык", "side": "Сторона", "name": "Игра",
                              "reviews_total": "Отзывов", "crossover_reviewers": "Людей со стыка в выборке",
                              "crossover_est_full": "Оценка людей со стыка"}).drop(columns=["appid"])
        write(w, "Игры для поиска", t, {"Стык": 20, "Игра": 36})
    print("written", DATA / "audience_and_creators.xlsx")


if __name__ == "__main__":
    main()
