"""Step 6: Excel workbook for manual review of the collected games.

Output: data/games_review.xlsx
"""
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import DATA

COLS = {
    "name": "Игра",
    "url": "Ссылка",
    "release_text": "Дата выхода",
    "price_usd": "Цена, $",
    "reviews_total": "Отзывы",
    "positive_pct": "% положит.",
    "sales_est_low": "Продажи ≈ от",
    "sales_est_high": "Продажи ≈ до",
    "developers": "Разработчик",
    "publishers": "Издатель",
    "big_publisher_match": "Крупный AA-издатель",
    "early_access": "Ранний доступ",
    "dog_rank": "Ранг тега Dogs",
    "hunting_rank": "Ранг тега Hunting",
    "tags": "Теги (по убыванию голосов)",
    "short_description": "Описание",
}
FLAG_COLS = {f"el_{e}": n for e, n in [("dog", "Собака"), ("hunting", "Охота"),
                                        ("friendslop", "Френдслоп"), ("horror", "Хоррор"),
                                        ("roguelite", "Рогалик"), ("extraction", "Экстракшн")]}
WIDTHS = {"Игра": 34, "Ссылка": 12, "Теги (по убыванию голосов)": 70, "Описание": 80,
          "Разработчик": 24, "Издатель": 24}

LEGEND = [
    ("Что это", "Игры Steam, собранные по тегам элементов игры (см. docs/methodology.md)."),
    ("Отзывы", "Все отзывы Steam на всех языках."),
    ("Продажи ≈", "Грубая оценка: отзывы × 20 … × 60. Точность ~2 раза в обе стороны."),
    ("Ранг тега", "Место тега среди тегов игры: 1 = тег, за который проголосовало больше всего игроков. "
                  "Ранг 15–20 обычно значит, что элемент в игре второстепенный."),
    ("Френдслоп", "Рабочее определение: Online Co-Op + (Funny/Comedy/Memes/Dark Comedy/Physics/Party Game), не MMO."),
    ("Собака", "Тег Dogs или слово dog/puppy/hound/corgi… в названии."),
    ("AAA исключены", "Игры крупнейших издателей — отдельный лист, в анализ не входят."),
    ("Крупный AA-издатель", "Помечены, но оставлены в анализе."),
]


def sheet(writer, name, df):
    cols = list(COLS) + list(FLAG_COLS)
    out = df[cols].rename(columns={**COLS, **FLAG_COLS})
    for c in FLAG_COLS.values():
        out[c] = out[c].map({True: "да", False: ""})
    out["Ранний доступ"] = out["Ранний доступ"].map({True: "да", False: ""})
    out.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(out.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 13)
        ws.cell(1, i).font = Font(bold=True)
    link_col = list(out.columns).index("Ссылка") + 1
    for r in range(2, len(out) + 2):
        cell = ws.cell(r, link_col)
        cell.hyperlink, cell.value, cell.style = cell.value, "Steam", "Hyperlink"


def main():
    df = pd.read_csv(DATA / "games.csv")
    for c in [c for c in df.columns if c.startswith("el_")] + ["early_access", "coming_soon", "is_aaa", "released", "serious_indie"]:
        df[c] = df[c].astype(bool)
    df["big_publisher_match"] = df.big_publisher_match.fillna("")
    released = df[df.released & ~df.is_aaa]
    with pd.ExcelWriter(DATA / "games_review.xlsx", engine="openpyxl") as w:
        pd.DataFrame(LEGEND, columns=["Поле", "Пояснение"]).to_excel(w, sheet_name="Пояснения", index=False)
        w.sheets["Пояснения"].column_dimensions["A"].width = 22
        w.sheets["Пояснения"].column_dimensions["B"].width = 120
        sheet(w, "Френдслоп", released[released.el_friendslop])
        sheet(w, "Собака", released[released.el_dog])
        sheet(w, "Охота", released[released.el_hunting])
        sheet(w, "Экстракшн", released[released.el_extraction])
        sheet(w, "Скоро выйдут", df[df.coming_soon & ~df.is_aaa])
        sheet(w, "AAA исключены", df[df.is_aaa])
        sheet(w, "Все вышедшие", released)
    print("written", DATA / "games_review.xlsx")


if __name__ == "__main__":
    main()
