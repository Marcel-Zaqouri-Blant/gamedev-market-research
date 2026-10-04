"""Step 13: Excel workbook with the profile results (what people at each intersection play).

Reads data/profiles_element_shares.csv and data/profiles_top_games.csv,
writes data/profiles.xlsx.
"""
import math

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import DATA

EL = ["dog", "hunting", "friendslop", "horror", "roguelite", "extraction"]
RU = {"dog": "собака", "hunting": "охота", "friendslop": "френдслоп",
      "horror": "хоррор", "roguelite": "рогалик", "extraction": "экстракшн"}


def ru_group(g):
    if g.startswith("control:"):
        return "все игроки: " + RU[g.split(":")[1]]
    return " + ".join(RU[x] for x in g.split("+"))


def write(w, name, df, widths=None):
    df.to_excel(w, sheet_name=name, index=False)
    ws = w.sheets[name]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, 14)
        ws.cell(1, i).font = Font(bold=True)


def main():
    sh = pd.read_csv(DATA / "profiles_element_shares.csv").set_index("group")
    top = pd.read_csv(DATA / "profiles_top_games.csv")

    people = pd.DataFrame({"Группа": [ru_group(g) for g in sh.index],
                           "Открытых профилей": sh.public_with_reviews.values,
                           "Других отзывов на человека": sh.avg_other_reviews.values})
    for e in EL:
        people[f"Играли в {RU[e]}, % людей"] = sh[f"people_{e}"].values
    shares = pd.DataFrame({"Группа": [ru_group(g) for g in sh.index]})
    for e in EL:
        shares[f"{RU[e]}, % отзывов"] = sh[f"share_{e}"].values

    # per pair: compare with the two controls
    rows = []
    for g in sh.index:
        if g.startswith("control:"):
            continue
        a, b = g.split("+")
        ca, cb = sh.loc[f"control:{a}"], sh.loc[f"control:{b}"]
        # distance between review compositions, in relative terms (mean |log ratio|):
        # people at a stык review twice as much as controls, so % of people would
        # be inflated by activity; review shares are not, and log ratios keep small
        # elements (dog, hunting) from being drowned out by horror
        vec = lambda r: [max(r[f"share_{e}"], 0.05) for e in EL]
        dist = lambda x, y: sum(abs(math.log(p / q)) for p, q in zip(vec(x), vec(y))) / len(EL)
        da, db = dist(sh.loc[g], ca), dist(sh.loc[g], cb)
        row = {"Стык": ru_group(g), "Людей": int(sh.loc[g].public_with_reviews),
               "Других отзывов на человека": sh.loc[g].avg_other_reviews}
        for e in EL:
            row[f"{RU[e]}, % их отзывов"] = sh.loc[g][f"share_{e}"]
        row.update({"Элемент 1": RU[a], "Отличие от всех игроков 1": round(da, 2),
                    "Элемент 2": RU[b], "Отличие от всех игроков 2": round(db, 2),
                    "Ближе к": RU[a] if da < db else RU[b]})
        rows.append(row)
    comp = pd.DataFrame(rows)

    legend = pd.DataFrame([
        ("Что это", "Публичные профили Steam: до 50 последних отзывов каждого человека. "
                    "Отзывы на игры, через которые человек попал в выборку, исключены."),
        ("Стык", "120 случайных людей, писавших отзывы на игры обоих элементов (игры-гибриды не в счёт)."),
        ("Все игроки: X", "Контрольная группа: 120 случайных людей, писавших отзывы на игры элемента X."),
        ("Играли ещё в X, %", "Доля людей группы, у которых в профиле есть хотя бы одна другая игра с элементом X."),
        ("% отзывов", "Какая часть всех их других отзывов приходится на игры с элементом."),
        ("Отличие от всех игроков", "Насколько состав отзывов стыка (колонки «% их отзывов») отличается "
                                    "от контрольной группы элемента, в относительных долях: меньше = похожее. "
                                    "«Ближе к» — на чью аудиторию стык похож больше."),
        ("Почему не % людей", "Люди со стыка пишут вдвое больше отзывов, чем контрольные группы, поэтому "
                              "«% людей, игравших в X» у них выше по всем жанрам сразу. Для сравнения "
                              "групп используется состав отзывов, он от активности не зависит."),
        ("Lift", "Во сколько раз чаще эту игру обозревают люди группы, чем все контрольные группы вместе. "
                 "В списке — только игры, которые обозревали 8+ человек группы."),
    ], columns=["Что", "Пояснение"])

    # 120 people per group: below 8 reviewers a game is noise (one curator, one friend group)
    t = top[top.people >= 8].copy()
    t["elements"] = t.elements.fillna("—")
    t["group"] = t.group.map(ru_group)
    t = t.rename(columns={"group": "Группа", "name": "Игра", "people": "Людей",
                          "pct_of_group": "% группы", "pct_of_controls": "% контрольных",
                          "lift_vs_controls": "Lift", "elements": "Элементы"}).drop(columns=["appid"])

    with pd.ExcelWriter(DATA / "profiles.xlsx", engine="openpyxl") as w:
        write(w, "Пояснения", legend, {"Что": 18, "Пояснение": 130})
        write(w, "Кто на стыке", comp, {"Стык": 22})
        write(w, "Во что играют (люди)", people, {"Группа": 26})
        write(w, "Во что играют (отзывы)", shares, {"Группа": 26})
        write(w, "Частые игры", t, {"Группа": 22, "Игра": 40, "Элементы": 26})
    print("written", DATA / "profiles.xlsx")


if __name__ == "__main__":
    main()
